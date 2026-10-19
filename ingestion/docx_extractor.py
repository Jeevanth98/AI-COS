"""
AI-COS — DOCX Text Extractor
Uses python-docx to extract paragraphs, tables, and headings from DOCX files.
"""

from __future__ import annotations
import logging
from pathlib import Path
from typing import Dict, Any, List

logger = logging.getLogger(__name__)


def extract_docx(file_path: str | Path) -> Dict[str, Any]:
    """
    Extract text and structural information from a DOCX file.

    Returns
    -------
    dict with keys:
        full_text   : str
        paragraphs  : list[{text, style, is_heading, heading_level}]
        tables      : list[{table_index, rows: list[list[str]]}]
        total_chars : int
        warnings    : list[str]
        error       : str | None
    """
    file_path = Path(file_path)
    result: Dict[str, Any] = {
        "full_text": "",
        "paragraphs": [],
        "tables": [],
        "total_chars": 0,
        "warnings": [],
        "error": None,
    }

    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError:
        result["error"] = "python-docx is not installed."
        return result

    if not file_path.exists():
        result["error"] = f"File not found: {file_path}"
        return result

    try:
        doc = Document(str(file_path))
    except Exception as exc:
        result["error"] = f"Cannot open DOCX: {exc}"
        return result

    paragraphs: List[Dict] = []
    text_parts: List[str] = []

    # ── Extract paragraphs (preserving heading hierarchy) ─────────────────
    for para in doc.paragraphs:
        text = para.text.strip()
        style_name = para.style.name if para.style else ""
        is_heading = "Heading" in style_name
        heading_level = None
        if is_heading:
            # e.g. "Heading 1" → 1
            parts = style_name.split()
            if len(parts) >= 2 and parts[-1].isdigit():
                heading_level = int(parts[-1])

        paragraphs.append({
            "text": text,
            "style": style_name,
            "is_heading": is_heading,
            "heading_level": heading_level,
        })
        if text:
            text_parts.append(text)

    # ── Extract tables ─────────────────────────────────────────────────────
    tables: List[Dict] = []
    for t_idx, table in enumerate(doc.tables):
        rows: List[List[str]] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)
            # Also add table content to full text
            text_parts.append(" | ".join(cells))
        tables.append({"table_index": t_idx, "rows": rows})

    full_text = "\n".join(text_parts)
    result["full_text"] = full_text
    result["paragraphs"] = paragraphs
    result["tables"] = tables
    result["total_chars"] = len(full_text.strip())

    if result["total_chars"] == 0:
        result["warnings"].append("No text was extracted from this DOCX file.")

    logger.debug("DOCX extracted: %s  chars=%d  paragraphs=%d  tables=%d",
                 file_path.name, result["total_chars"], len(paragraphs), len(tables))
    return result
